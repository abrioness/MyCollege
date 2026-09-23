# -*- coding: utf-8 -*-
"""Genera el Manual de Usuario en Word del sistema WebColegio."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Pt, RGBColor, Emu

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent / "Manual_Usuario_WebColegio.docx"
LOGO = ROOT / "wwwroot" / "Image" / "colegiosanfrancisco.png"

NAVY = RGBColor(0x1B, 0x36, 0x5D)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
DARK = RGBColor(0x2C, 0x2C, 0x2C)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = "F4F7FB"
HEADER_BG = "1B365D"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=DARK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "C5CDD8")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Colegio Parroquial San Francisco Javier  ·  Manual de usuario  ·  Página ")
    set_run_font(run, size=9, color=GRAY)
    add_page_number(p)
    run2 = p.add_run("")
    set_run_font(run2, size=9, color=GRAY)


def add_header(section, text):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    set_run_font(run, size=9, italic=True, color=NAVY)


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=16, bold=True, color=NAVY)
    # underline bar
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "1B365D")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=13, bold=True, color=NAVY)
    return p


def h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True, color=RGBColor(0x2E, 0x5A, 0x88))
    return p


def para(doc, text, *, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def bullets(doc, items, numbered=False):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.space_after = Pt(3)
        prefix = f"{i}. " if numbered else "• "
        run = p.add_run(prefix + item)
        set_run_font(run, size=11)
    return


def note(doc, text, title="Importante"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "FFF6E0")
    set_cell_border(cell)
    p = cell.paragraphs[0]
    r1 = p.add_run(title + ". ")
    set_run_font(r1, size=11, bold=True, color=GOLD)
    r2 = p.add_run(text)
    set_run_font(r2, size=11)
    doc.add_paragraph()


def tip(doc, text, title="Consejo"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "E8F1FB")
    set_cell_border(cell)
    p = cell.paragraphs[0]
    r1 = p.add_run(title + ". ")
    set_run_font(r1, size=11, bold=True, color=NAVY)
    r2 = p.add_run(text)
    set_run_font(r2, size=11)
    doc.add_paragraph()


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, HEADER_BG)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=WHITE)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            if r_i % 2 == 1:
                shade_cell(cell, ROW_ALT)
            set_cell_border(cell)
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10)
    doc.add_paragraph()


def toc_line(doc, num, title, page=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"{num}   {title}")
    set_run_font(run, size=12, color=NAVY)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    add_header(section, "Sistema de gestión escolar")
    add_footer(section)

    # -------- PORTADA --------
    for _ in range(3):
        doc.add_paragraph()
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(LOGO), width=Cm(4.2))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run("COLEGIO PARROQUIAL")
    set_run_font(run, size=14, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("SAN FRANCISCO JAVIER")
    set_run_font(run, size=22, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("En todo amar y servir")
    set_run_font(run, size=12, italic=True, color=GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    run = p.add_run("MANUAL DE USUARIO")
    set_run_font(run, size=26, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Sistema de gestión escolar WebColegio")
    set_run_font(run, size=14, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    run = p.add_run("Septiembre 2026")
    set_run_font(run, size=12, color=DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Uso interno del personal del colegio")
    set_run_font(run, size=11, italic=True, color=GRAY)

    doc.add_page_break()

    # -------- ÍNDICE --------
    h1(doc, "Índice")
    contents = [
        ("1.", "Presentación del sistema"),
        ("2.", "Acceso e inicio de sesión"),
        ("3.", "Perfiles y menú de trabajo"),
        ("4.", "Uso general de pantallas y reportes"),
        ("5.", "Alumnos"),
        ("6.", "Matrícula"),
        ("7.", "Traslado entre colegios"),
        ("8.", "Notas"),
        ("9.", "Inventario"),
        ("10.", "Caja"),
        ("11.", "Arqueo diario y cierre de caja"),
        ("12.", "Avisos por correo"),
        ("13.", "Avisos WhatsApp"),
        ("14.", "Reportes"),
        ("15.", "Tarifas y ciclo lectivo"),
        ("16.", "Usuarios"),
        ("17.", "Preguntas frecuentes y recomendaciones"),
    ]
    for n, t in contents:
        toc_line(doc, n, t)

    doc.add_page_break()

    # 1
    h1(doc, "1. Presentación del sistema")
    para(
        doc,
        "Este manual explica cómo usar el sistema de gestión escolar del Colegio Parroquial "
        "San Francisco Javier. Está pensado para el personal de caja, secretaría, docentes "
        "y administración. No es un documento técnico: describe lo que se ve en pantalla y "
        "los pasos del día a día.",
    )
    para(doc, "Con el sistema se puede:")
    bullets(
        doc,
        [
            "Registrar y consultar estudiantes.",
            "Matricular y dar seguimiento al ciclo lectivo.",
            "Registrar notas.",
            "Cobrar matrícula, mensualidades, otros ingresos y egresos.",
            "Consultar el estado de cuenta y la mora.",
            "Cerrar caja y generar el arqueo del día.",
            "Avisar a tutores por correo o WhatsApp cuando hay saldo pendiente.",
            "Consultar reportes y administrar usuarios y tarifas (según el perfil).",
        ],
    )
    note(
        doc,
        "Cada persona solo ve los menús de su rol. Si no aparece una opción, es normal: "
        "pida al administrador que revise el perfil o los permisos.",
    )

    # 2
    h1(doc, "2. Acceso e inicio de sesión")
    h2(doc, "2.1 Entrar al sistema")
    bullets(
        doc,
        [
            "Abra el navegador (se recomienda Google Chrome o Microsoft Edge).",
            "Ingrese a la dirección del sistema que le indicó administración.",
            "En Usuario escriba su nombre de usuario.",
            "En Contraseña escriba su clave. Distingue mayúsculas y minúsculas.",
            "Pulse Iniciar sesión.",
        ],
        numbered=True,
    )
    para(
        doc,
        "Si los datos son correctos, verá el menú izquierdo y la pantalla de inicio de su perfil "
        "(estado de cuenta para caja y administración; notas para docentes).",
    )
    h2(doc, "2.2 Si no puede entrar")
    bullets(
        doc,
        [
            "Revise que no tenga mayúsculas activadas y que el usuario esté escrito igual que se lo asignaron.",
            "Si olvida la contraseña, solicite el restablecimiento a administración. No comparta su clave.",
            "Si el sistema indica que el servicio no está disponible, espere unos minutos e intente de nuevo. Si continúa, avise a administración.",
            "No use la cuenta de otra persona: cada recibo y cada movimiento quedan asociados a quien inició sesión.",
        ],
    )
    h2(doc, "2.3 Cerrar sesión")
    para(
        doc,
        "Al terminar el turno, cierre sesión desde el menú de usuario (parte superior). "
        "No deje la sesión abierta en una computadora compartida.",
    )

    # 3
    h1(doc, "3. Perfiles y menú de trabajo")
    para(
        doc,
        "El menú izquierdo cambia según el rol. Estos son los perfiles habituales y lo que cada uno usa:",
    )
    add_table(
        doc,
        ["Perfil", "Para qué se usa principalmente"],
        [
            (
                "Cajero",
                "Alumnos, matrícula, caja, estado de cuenta, avisos, reportes de caja, arqueo y cierre.",
            ),
            (
                "Secretaria",
                "Notas, reporte de alumnos, inventario y arqueo (consulta o generación, según se habilite).",
            ),
            (
                "Docente",
                "Registro de notas y consulta de alumnos.",
            ),
            (
                "Admin / UserSystem",
                "Todo lo anterior, más tarifas, ciclo lectivo, usuarios, anulación de recibos y permisos.",
            ),
            (
                "Director",
                "Consulta de reportes académicos (alumnos y notas), según el menú asignado.",
            ),
        ],
    )
    h2(doc, "3.1 Módulos del menú")
    add_table(
        doc,
        ["Módulo", "Opciones"],
        [
            ("Alumnos", "Agregar estudiante. Traslado entre colegios."),
            ("Matrícula", "Matricular y ver el reporte de matrícula."),
            ("Notas", "Agregar notas y reporte de alumnos."),
            ("Inventario", "Agregar productos o artículos."),
            (
                "Caja",
                "Recibo de mensualidad, recibo varios, egresos, traslado, estado de cuenta, cerrar caja, avisos por correo y WhatsApp.",
            ),
            (
                "Reportes",
                "Alumnos, notas, inventario, caja mensualidad, caja varios, egresos y estado de cuenta.",
            ),
            ("Arqueo diario", "Lista de arqueos y generar arqueo del día."),
            ("Tarifas y ciclo", "Costos de matrícula y mensualidad; ciclo actual y siguiente."),
            ("Usuario", "Lista de usuarios, alta de usuario y permisos de menú."),
        ],
    )

    # 4
    h1(doc, "4. Uso general de pantallas y reportes")
    bullets(
        doc,
        [
            "Las listas tienen buscador y se pueden ordenar al pulsar el título de la columna.",
            "En varios reportes hay exportación a Excel. Use esa opción para llevar el listado a una hoja de cálculo.",
            "Los mensajes en verde confirman que la operación se guardó. Los mensajes en amarillo o rojo piden corregir un dato o avisan que algo no se pudo completar.",
            "Si una pantalla queda vacía, primero revise filtros (ciclo, recinto, fechas o búsqueda). No siempre es un error.",
            "El número de recibo lo asigna el sistema. No lo invente ni lo cambie a mano.",
        ],
    )
    tip(
        doc,
        "Antes de guardar un recibo, revise alumno, concepto, ciclo, recinto, fecha de pago, meses y monto. "
        "Corregir después implica anular (solo administración) y volver a emitir.",
    )

    # 5
    h1(doc, "5. Alumnos")
    h2(doc, "5.1 Agregar estudiante")
    para(doc, "Ruta: Alumnos → Agregar estudiante.")
    para(
        doc,
        "El código del estudiante lo genera el sistema. Complete los datos personales, académicos y de familia. "
        "Entre los campos más usados están:",
    )
    bullets(
        doc,
        [
            "Nombre y apellido.",
            "Código MINED y código de persona, si aplica.",
            "Fecha de nacimiento, sexo, dirección, barrio, departamento y municipio.",
            "Grado, turno, modalidad, recinto y ciclo.",
            "Datos de madre, padre y tutor: nombre, cédula y teléfono.",
            "Correo (se usa para avisos de saldo). Si el alumno no tiene correo, se puede usar el del usuario tutor.",
            "Beca completa o media beca, si corresponde.",
            "Observaciones.",
        ],
    )
    para(
        doc,
        "El teléfono del tutor (o de madre/padre) es el que se usa para WhatsApp. "
        "Conviene escribirlo completo y sin letras, por ejemplo 88881234.",
    )
    h2(doc, "5.2 Reporte de alumnos")
    para(
        doc,
        "Ruta: Reportes → Reporte de alumnos (o Notas → Reporte de alumnos, según el perfil).",
    )
    para(
        doc,
        "Ahí se busca, se edita la ficha y, si es administrador, se puede anular un registro. "
        "Anular un estudiante lo deja fuera de las operaciones activas; no lo use para un error menor de captura: edite primero.",
    )

    # 6
    h1(doc, "6. Matrícula")
    para(doc, "Ruta: Matrícula → Matrícula (para registrar) o Reporte de matrícula (para consultar).")
    bullets(
        doc,
        [
            "Busque al alumno y confirme que la ficha esté activa.",
            "Elija ciclo, recinto, modalidad, grado y el resto de datos académicos.",
            "Guarde. Puede imprimir la constancia o ficha de matrícula desde la opción Imprimir.",
            "En el reporte puede filtrar por ciclo, estado y texto de búsqueda.",
        ],
        numbered=True,
    )
    note(
        doc,
        "La matrícula académica y el cobro de matrícula son pasos distintos. Primero se registra al alumno "
        "y su matrícula; después se cobra en Caja → Recibo de caja mensualidad, con el concepto de matrícula "
        "o matrícula completa.",
    )
    para(
        doc,
        "Si el alumno se matricula tarde (por ejemplo en marzo o después), el sistema no le cobra mensualidades "
        "de los meses anteriores a su ingreso. Empieza a deber desde el mes de esa matrícula.",
    )

    # 7
    h1(doc, "7. Traslado entre colegios")
    para(doc, "Ruta: Alumnos → Traslado entre colegios, o Caja → Traslado entre colegios.")
    para(
        doc,
        "Se usa cuando el estudiante cambia de recinto o colegio dentro del mismo ciclo. "
        "No se da de baja la ficha ni se crea otra matrícula: se actualiza el recinto destino "
        "y se cobran la matrícula y las mensualidades del colegio de llegada desde el mes de ingreso.",
    )
    bullets(
        doc,
        [
            "Busque al alumno y confirme el recinto de origen.",
            "Indique el recinto destino y la fecha real del traslado (esa fecha define el mes de ingreso).",
            "Revise el resumen: qué se cobra en el destino y desde qué mes.",
            "Guarde el traslado y genere el recibo de cobro del destino.",
        ],
        numbered=True,
    )
    note(
        doc,
        "Los pagos hechos en el colegio de origen no se aplican al destino. "
        "Ejemplo: si Mariel se trasladó el 23 de febrero, en San Francisco debe febrero y los meses siguientes; "
        "no se le cobran meses anteriores al traslado. Si paga el 25 de marzo solo febrero, la mora es de febrero. "
        "Si paga el 2 de abril febrero y marzo, la mora es de esos dos meses (C$ 20).",
    )

    # 8
    h1(doc, "8. Notas")
    para(doc, "Ruta: Notas → Agregar notas.")
    bullets(
        doc,
        [
            "Busque al alumno. El sistema carga modalidad, nivel y recinto si la matrícula corresponde al período lectivo actual.",
            "Confirme o corrija recinto, asignatura, tipo de evaluación y período de evaluación.",
            "Registre las calificaciones y guarde.",
            "El reporte de notas (Reportes → Reporte de notas) sirve para consultar e imprimir.",
        ],
        numbered=True,
    )
    para(
        doc,
        "Si al buscar un alumno no cargan modalidad o nivel, revise que tenga matrícula activa en el ciclo vigente.",
    )

    # 9
    h1(doc, "9. Inventario")
    para(doc, "Ruta: Inventario → Agregar inventario, y Reportes → Reporte de inventario.")
    para(
        doc,
        "Registre el artículo (nombre, código, categoría y existencias). "
        "El reporte permite consultar el catálogo. Administración puede anular un producto; "
        "queda desactivado y deja de mostrarse en el listado activo.",
    )

    # 10
    h1(doc, "10. Caja")
    para(
        doc,
        "Caja concentra los cobros del colegio. Hay tres tipos de recibo y varias consultas. "
        "Siempre confirme recinto, ciclo y fecha antes de guardar.",
    )

    h2(doc, "10.1 Recibo de caja mensualidad")
    para(doc, "Ruta: Caja → Recibo de caja mensualidad.")
    para(doc, "Pasos habituales:")
    bullets(
        doc,
        [
            "El número de recibo aparece solo (no se edita).",
            "Busque al alumno por nombre. Al elegirlo se cargan recinto, modalidad, grado y datos del ciclo.",
            "Elija el concepto: matrícula, matrícula completa, mensualidad, rifa, promoción u otro que esté habilitado.",
            "Confirme tipo de recibo, ciclo escolar, centro de estudio, modalidad, nivel y método de pago.",
            "Indique la fecha de pago (fecha real en que el tutor cancela, no necesariamente el día de hoy si está regularizando un cobro físico).",
            "En mensualidad, pulse Seleccionar meses y marque solo los meses que está pagando o abonando.",
            "Escriba el monto. Puede ser el mes completo, un abono menor o un excedente (el sobrante pasa al mes siguiente).",
            "Revise la mora si el sistema la muestra. Es de solo lectura.",
            "Agregue observaciones si hace falta y guarde. Imprima el recibo y entréguelo al tutor.",
        ],
        numbered=True,
    )

    h3(doc, "Matrícula, matrícula completa y continuidad")
    bullets(
        doc,
        [
            "La matrícula debe estar cancelada o abonada antes de cobrar mensualidades sueltas.",
            "En matrícula completa, la mensualidad incluida se registra siempre como enero, aunque la fecha del recibo sea de otro mes.",
            "La casilla Es continuidad se usa cuando el alumno sigue activo en este ciclo aunque no haya pagado matrícula ni mensualidades. El pago cubre matrícula más enero de este ciclo (no del siguiente). Después se pueden cobrar los meses anteriores al mes de esa matrícula.",
            "Mensualidad, morosos y adelantos usan el ciclo actual. El siguiente ciclo solo aplica a reserva o matrícula nueva (normalmente desde octubre).",
        ],
    )

    h3(doc, "Abonos")
    para(
        doc,
        "Si el tutor no cubre el mes completo, registre el monto entregado. El sistema deja el saldo del mes. "
        "Si entrega de más, el excedente se aplica al mes siguiente.",
    )

    h2(doc, "10.2 Mora")
    para(
        doc,
        "La mora es C$ 10 por cada mes vencido respecto a la fecha de pago del recibo. "
        "No se cobra el mes en curso de esa fecha. El cálculo no usa “hoy” del calendario si usted "
        "está registrando un pago con otra fecha de emisión.",
    )
    add_table(
        doc,
        ["Situación", "Cómo se cobra la mora"],
        [
            (
                "Alumno regular del ciclo",
                "C$ 10 por cada mes seleccionado que ya venció antes del mes de la fecha de pago.",
            ),
            (
                "Traslado (ej. 23 de febrero)",
                "Empieza en el mes del traslado. No se cobran meses anteriores al ingreso en el destino.",
            ),
            (
                "Matrícula tardía (marzo o después)",
                "Empieza en el mes de esa matrícula. No se cobran meses anteriores.",
            ),
            (
                "Pago el 25/03 de solo febrero",
                "Mora de febrero: C$ 10.",
            ),
            (
                "Pago el 02/04 de febrero y marzo",
                "Mora de febrero y marzo: C$ 20.",
            ),
            (
                "Recibo anulado",
                "No cuenta como pago. El mes vuelve a quedar pendiente y entra de nuevo a la mora si aplica.",
            ),
        ],
    )
    note(
        doc,
        "La mora la calcula el sistema. No la escriba a mano. Si no aparece, revise fecha de pago, "
        "meses seleccionados y que el alumno tenga matrícula o traslado bien registrado.",
    )

    h2(doc, "10.3 Recibo de caja varios")
    para(doc, "Ruta: Caja → Recibo de caja varios.")
    para(
        doc,
        "Se usa para cobros que no son la mensualidad académica habitual (útiles, actividades, otros ingresos). "
        "Busque al alumno o escriba de quién se recibe el dinero, elija concepto, método de pago, monto y fecha, y guarde.",
    )

    h2(doc, "10.4 Recibo de caja egresos")
    para(doc, "Ruta: Caja → Recibo de caja egresos.")
    para(
        doc,
        "Registra salidas de dinero (compras, reembolsos u otros egresos autorizados). "
        "Complete a quién se paga, el concepto, el monto y la observación. "
        "Estos movimientos también entran al arqueo del día.",
    )

    h2(doc, "10.5 Estado de cuenta")
    para(doc, "Ruta: Caja → Estado de cuenta, o Reportes → Estado de cuenta.")
    para(
        doc,
        "Muestra, por alumno, lo pagado y lo pendiente de matrícula, mensualidades y otros conceptos del ciclo. "
        "Sirve para atender al tutor y para las campañas de aviso. "
        "Los estudiantes retirados o inactivos se distinguen para no mezclarlos con la cartera activa.",
    )

    h2(doc, "10.6 Anular un recibo")
    para(
        doc,
        "Solo administración puede anular. Se hace desde el reporte correspondiente "
        "(caja mensualidad, caja varios o egreso), con el botón Anular.",
    )
    bullets(
        doc,
        [
            "El recibo anulado sigue visible en reportes, para control.",
            "No se toma en cuenta en ninguna validación de saldo, mora ni meses pagados.",
            "Si anula un recibo de mensualidad con varias líneas (varios meses), se anulan todas las del mismo número.",
            "Después de anular, si el tutor ya pagó de nuevo, emita un recibo nuevo con la fecha y los meses correctos.",
        ],
    )

    # 11
    h1(doc, "11. Arqueo diario y cierre de caja")
    h2(doc, "11.1 Cerrar caja")
    para(doc, "Ruta: Caja → Cerrar caja.")
    para(
        doc,
        "Al final del turno, cierre caja. Ese cierre fija el corte del día: el arqueo mostrará "
        "los cobros hasta esa hora. Lo que se cobre después entra al siguiente cierre.",
    )
    h2(doc, "11.2 Generar arqueo")
    para(doc, "Ruta: Arqueo diario → Generar arqueo (se abre en una ventana para imprimir).")
    bullets(
        doc,
        [
            "El cajero ve el arqueo de su recinto de trabajo.",
            "Administración elige el recinto y pulsa Generar arqueo del recinto seleccionado.",
            "El documento muestra ingresos, egresos, denominaciones y el turno (hora de inicio y fin de la ventana de caja).",
            "Si la caja aún está abierta, el arqueo muestra cobros hasta el momento. Cierre caja para dejar el corte definitivo.",
            "Si el arqueo ya está cerrado, no lo vuelva a tratar como un segundo corte del mismo turno.",
        ],
    )
    para(doc, "La lista de arqueos (Arqueo diario → Lista de arqueos) sirve para consultar cortes anteriores y exportar.")

    # 12
    h1(doc, "12. Avisos por correo")
    para(doc, "Ruta: Caja → Avisos por correo.")
    para(
        doc,
        "Lista a los alumnos con saldo pendiente y el correo del tutor. "
        "El correo se toma de la ficha del alumno o, si no hay, del usuario tutor. "
        "El mensaje sale como no responder: el tutor no debe contestar ese aviso; la bandeja no se revisa.",
    )
    bullets(
        doc,
        [
            "Enviar: un alumno, los seleccionados o la campaña completa (todos los que tienen correo y saldo).",
            "Programar: elija una fecha. El envío masivo sale a partir de las 8:00 de ese día. Si programa hoy y ya pasó esa hora, se envía de una vez.",
            "Cancelar programación si se equivocó de fecha.",
            "Revise el historial de últimos envíos al pie de la pantalla.",
        ],
    )
    note(
        doc,
        "Si los botones de envío no aparecen, el correo del colegio aún no está configurado. "
        "Pida a administración que active la cuenta noreply. Mientras tanto puede revisar quién tiene correo y quién no, para completar fichas.",
    )

    # 13
    h1(doc, "13. Avisos WhatsApp")
    para(doc, "Ruta: Caja → Avisos WhatsApp.")
    para(
        doc,
        "Los avisos masivos deben salir del número oficial del colegio: +505 8425 8684. "
        "El tutor debe ver ese número, no el teléfono personal de quien está en caja.",
    )
    h2(doc, "13.1 Envío uno a uno (siempre disponible)")
    para(
        doc,
        "Use el ícono verde de WhatsApp en cada fila. Se abre el chat con el mensaje listo. "
        "Para que el tutor vea el 84258684, entre a WhatsApp Web con esa cuenta del colegio antes de enviar.",
    )
    h2(doc, "13.2 Envío masivo desde el 84258684")
    para(
        doc,
        "Cuando el número del colegio esté vinculado para envío automático, aparecerán los botones "
        "Enviar seleccionados, Campaña completa y la programación por fecha (igual que el correo, a partir de las 8:00).",
    )
    para(
        doc,
        "Mientras el número no esté vinculado en WhatsApp Business, el sistema no puede disparar "
        "cientos de mensajes solo. Tener el 84258684 no basta: hay que asociarlo a la cuenta institucional.",
    )
    tip(
        doc,
        "Antes de una campaña, complete teléfonos de tutor en las fichas. "
        "Un número de 8 dígitos de Nicaragua es suficiente; el sistema antepone el 505.",
    )

    # 14
    h1(doc, "14. Reportes")
    para(
        doc,
        "Los reportes están en el menú Reportes. Use filtros, buscador y, cuando exista, Excel.",
    )
    add_table(
        doc,
        ["Reporte", "Qué consulta"],
        [
            ("Alumnos", "Fichas, códigos, grado, recinto y estado."),
            ("Notas", "Calificaciones por alumno, asignatura y período."),
            ("Inventario", "Artículos y existencias."),
            ("Caja mensualidad", "Recibos de matrícula y mensualidades, incluidos anulados (marcados)."),
            ("Caja varios", "Otros ingresos de caja."),
            ("Caja egreso", "Salidas de dinero."),
            ("Estado de cuenta", "Saldos y meses pendientes por alumno."),
            ("Usuarios", "Cuentas del sistema, rol, correo y fecha de creación."),
            ("Matrícula", "Alumnos matriculados en el ciclo, con filtros de estado."),
        ],
    )
    para(
        doc,
        "En reportes de caja, un recibo anulado se conserva para auditoría. "
        "No lo sume otra vez como ingreso del día ni lo use para decir que el mes ya está pagado.",
    )

    # 15
    h1(doc, "15. Tarifas y ciclo lectivo")
    para(
        doc,
        "Ruta: Tarifas y ciclo (solo administración). Ahí se definen los montos de matrícula y mensualidad "
        "por recinto, modalidad, grado y ciclo.",
    )
    bullets(
        doc,
        [
            "El ciclo actual se usa para mensualidades.",
            "La matrícula del siguiente ciclo se abre en octubre, sin cambiar el ciclo vigente. Así no se mezclan los cobros de octubre a diciembre con el año nuevo.",
            "Se pueden agregar o editar tarifas de matrícula y de mensualidad.",
            "Habilitar el siguiente ciclo (y copiar tarifas, si se pide) lo hace solo administración, cuando corresponda.",
        ],
    )
    note(
        doc,
        "No cambie tarifas a mitad de un cobro del día sin avisar a caja. "
        "Un monto distinto en pantalla y en el recibo físico genera reclamos al tutor.",
    )

    # 16
    h1(doc, "16. Usuarios")
    para(doc, "Ruta: Usuario → Reporte de usuarios / Agregar usuario (administración).")
    para(doc, "Al crear un usuario se definen:")
    bullets(
        doc,
        [
            "Nombre de usuario y nombre completo.",
            "Cédula, correo y colegio o recinto de trabajo.",
            "Rol (cajero, docente, secretaria, administrador, etc.).",
            "Si aplica, el vínculo con un alumno (útil para tutores).",
        ],
    )
    para(
        doc,
        "El reporte muestra la fecha de creación del usuario, el estado (activo o no) y permite editar. "
        "No desactive una cuenta de caja en medio del turno.",
    )
    para(
        doc,
        "Permisos de menú permite mostrar u ocultar opciones por rol. "
        "Úselo con cuidado: un cajero sin Caja no podrá cobrar.",
    )

    # 17
    h1(doc, "17. Preguntas frecuentes y recomendaciones")
    h2(doc, "¿El tutor dice que ya pagó y el estado de cuenta sigue en mora?")
    para(
        doc,
        "Busque el recibo. Si está anulado, no cuenta. Si el pago es de otro recinto (traslado), "
        "no abona el destino. Si el recibo es de otro ciclo, no cubre el ciclo actual.",
    )
    h2(doc, "¿Puedo cobrar mensualidad sin matrícula?")
    para(
        doc,
        "No, salvo el caso de continuidad (casilla especial) o las reglas que indiquen administración. "
        "El sistema pide matrícula cancelada o abonada antes de la mensualidad suelta.",
    )
    h2(doc, "¿Qué fecha pongo en el recibo?")
    para(
        doc,
        "La fecha real en que el tutor pagó. Esa fecha mueve la mora. "
        "Si registra al día siguiente un pago de ayer, use la fecha de ayer.",
    )
    h2(doc, "¿Los recibos anulados aparecen en el arqueo como ingreso?")
    para(
        doc,
        "No deben alterar el saldo ni las validaciones. Siguen en reportes para control. "
        "Si tiene duda en un arqueo concreto, revise con administración el corte de esa fecha.",
    )
    h2(doc, "¿Por qué un alumno no sale en avisos de correo o WhatsApp?")
    bullets(
        doc,
        [
            "No tiene saldo pendiente en el ciclo.",
            "Está retirado, inactivo o el registro de traslado histórico no aplica a la cartera actual.",
            "No tiene correo (avisos por correo) o no tiene teléfono de tutor, madre, padre o alumno (WhatsApp).",
        ],
    )
    h2(doc, "Buenas prácticas")
    bullets(
        doc,
        [
            "Complete ficha del alumno: teléfono, correo y cédula del tutor evitan errores en cobro y avisos.",
            "No use el mismo recibo para mezclar conceptos distintos si el sistema pide recibos separados.",
            "Cierre caja al terminar el turno, aunque vaya a generar el arqueo más tarde.",
            "No anule por error de un córdoba: primero confirme con administración.",
            "Los avisos masivos de correo son noreply. Si el tutor quiere aclarar, que llame a caja; no que responda el correo.",
            "WhatsApp masivo solo desde el 84258684 institucional.",
        ],
    )

    h1(doc, "Control del documento")
    add_table(
        doc,
        ["Dato", "Valor"],
        [
            ("Colegio", "Colegio Parroquial San Francisco Javier"),
            ("Documento", "Manual de usuario del sistema de gestión escolar"),
            ("Público", "Personal interno (caja, secretaría, docentes, administración)"),
            ("Fecha", "Septiembre 2026"),
            ("Versión", "1.0"),
        ],
    )
    para(
        doc,
        "Si una pantalla cambia después de una actualización, use este manual como guía general "
        "y consulte a administración ante cualquier diferencia.",
        italic=True,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
